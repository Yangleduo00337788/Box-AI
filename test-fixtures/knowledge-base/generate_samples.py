#!/usr/bin/env python3
"""Generate sample files for Box knowledge base manual testing."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent

# Shared facts — use these in RAG test questions
FACTS = {
    "product_name": "Box AI",
    "company": "星云科技有限公司",
    "support_email": "support@xingyun-tech.com",
    "hotline": "400-880-2024",
    "refund_days": "7",
    "enterprise_price": "9999",
    "max_agents_free": "3",
    "kb_storage_gb_pro": "50",
    "sla_uptime": "99.9%",
}


def write_text_files() -> None:
    (OUT / "01-产品介绍.txt").write_text(
        f"""Box 知识库测试文档（纯文本）

产品名称：{FACTS['product_name']}
运营主体：{FACTS['company']}

【客服】
- 邮箱：{FACTS['support_email']}
- 热线：{FACTS['hotline']}

【套餐】
免费版最多创建 {FACTS['max_agents_free']} 个智能体。
专业版知识库存储上限为 {FACTS['kb_storage_gb_pro']} GB。

【售后】
自开通日起 {FACTS['refund_days']} 天内可申请无理由退款（企业版除外）。
""",
        encoding="utf-8",
    )

    (OUT / "02-使用指南.md").write_text(
        f"""# Box 知识库使用指南

> 本文档用于 **RAG 检索测试**，内容与同目录其他格式文件一致。

## 快速开始

1. 在控制台创建知识库
2. 上传 TXT、MD、PDF、DOCX 等文档
3. 等待状态变为 **READY** 后在「RAG 测试」中提问

## 常见问题

### 企业版价格是多少？

企业版年费为 **{FACTS['enterprise_price']} 元/年**，含专属 SLA。

### 服务可用性承诺？

我们承诺月度可用性不低于 **{FACTS['sla_uptime']}**。

### 如何联系技术支持？

请发送邮件至 `{FACTS['support_email']}` 或拨打 `{FACTS['hotline']}`。
""",
        encoding="utf-8",
    )

    (OUT / "03-配置说明.json").write_text(
        json.dumps(
            {
                "product": FACTS["product_name"],
                "vendor": FACTS["company"],
                "support": {
                    "email": FACTS["support_email"],
                    "phone": FACTS["hotline"],
                },
                "plans": {
                    "free": {"max_agents": int(FACTS["max_agents_free"])},
                    "pro": {"knowledge_storage_gb": int(FACTS["kb_storage_gb_pro"])},
                    "enterprise": {"annual_price_cny": int(FACTS["enterprise_price"])},
                },
                "policy": {"refund_within_days": int(FACTS["refund_days"])},
                "sla": {"monthly_uptime_percent": FACTS["sla_uptime"]},
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    with (OUT / "04-套餐对比.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["套餐", "智能体上限", "知识库存储(GB)", "年费(元)", "备注"])
        w.writerow(["免费版", FACTS["max_agents_free"], "5", "0", "适合个人试用"])
        w.writerow(["专业版", "20", FACTS["kb_storage_gb_pro"], "1999", "含 RAG 测试"])
        w.writerow(["企业版", "不限", "500", FACTS["enterprise_price"], f"SLA {FACTS['sla_uptime']}"])

    (OUT / "05-部署日志片段.log").write_text(
        f"""2026-09-19T10:00:01Z INFO  box-knowledge  document parsed file=sample.docx
2026-09-19T10:00:02Z INFO  box-knowledge  chunk count=12 kb_id=demo
2026-09-19T10:00:05Z INFO  embedding model=text-embedding-3-small
2026-09-19T10:00:08Z INFO  index ready elasticsearch index=box:knowledge
2026-09-19T10:00:09Z INFO  support contact registered email={FACTS['support_email']}
""",
        encoding="utf-8",
    )


def write_docx() -> None:
    doc = Document()
    doc.add_heading("Box 企业采购说明（Word）", level=1)
    doc.add_paragraph(
        f"采购方在签署合同后，可通过 {FACTS['support_email']} 申请开通企业版。"
    )
    doc.add_paragraph(
        f"标准企业版年费为 {FACTS['enterprise_price']} 元人民币，包含 {FACTS['sla_uptime']} 可用性 SLA。"
    )
    doc.add_paragraph(
        f"非企业版用户可在 {FACTS['refund_days']} 天内申请退款；企业版按合同约定执行。"
    )
    doc.add_paragraph(f"技术支持热线：{FACTS['hotline']}。")
    doc.save(OUT / "06-企业采购说明.docx")


def write_xlsx() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "配额"
    ws.append(["项目", "数值", "说明"])
    ws.append(["产品", FACTS["product_name"], "Box 平台"])
    ws.append(["免费版智能体数", FACTS["max_agents_free"], "上限"])
    ws.append(["专业版知识库 GB", FACTS["kb_storage_gb_pro"], "存储配额"])
    ws.append(["企业版年费", FACTS["enterprise_price"], "元/年"])
    ws.append(["退款窗口", FACTS["refund_days"], "天"])
    ws.append(["客服邮箱", FACTS["support_email"], ""])
    wb.save(OUT / "07-配额表.xlsx")


def write_pptx() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(1))
    title.text_frame.text = f"{FACTS['product_name']} 培训幻灯片"
    title.text_frame.paragraphs[0].font.size = Pt(28)
    body = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(4))
    tf = body.text_frame
    tf.text = "知识库 RAG 测试要点"
    for line in [
        f"公司：{FACTS['company']}",
        f"企业版价格：{FACTS['enterprise_price']} 元/年",
        f"专业版知识库：{FACTS['kb_storage_gb_pro']} GB",
        f"支持邮箱：{FACTS['support_email']}",
        f"SLA：{FACTS['sla_uptime']}",
    ]:
        p = tf.add_paragraph()
        p.text = line
        p.level = 0
    prs.save(OUT / "08-培训要点.pptx")


def write_pdf() -> None:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen import canvas
    except ImportError:
        (OUT / "09-服务条款.pdf").write_bytes(b"%PDF-1.4\n% Box test placeholder\n")
        print("reportlab not installed; wrote minimal PDF placeholder")
        return

    font_path = Path("C:/Windows/Fonts/msyh.ttc")
    if font_path.exists():
        pdfmetrics.registerFont(TTFont("YaHei", str(font_path), subfontIndex=0))
        font = "YaHei"
    else:
        font = "Helvetica"

    pdf_path = OUT / "09-服务条款.pdf"
    c = canvas.Canvas(str(pdf_path), pagesize=A4)
    c.setFont(font, 14)
    y = 800
    lines = [
        "Box 服务条款摘要（PDF 测试）",
        f"产品：{FACTS['product_name']} | 主体：{FACTS['company']}",
        f"企业版年费 {FACTS['enterprise_price']} 元，SLA {FACTS['sla_uptime']}。",
        f"退款：开通后 {FACTS['refund_days']} 天内（企业版除外）。",
        f"联系：{FACTS['support_email']} / {FACTS['hotline']}",
    ]
    for line in lines:
        c.drawString(50, y, line)
        y -= 28
    c.save()


def write_readme() -> None:
    (OUT / "README.md").write_text(
        f"""# 知识库测试样例文件

本目录由 `generate_samples.py` 生成，各文件包含**同一套事实**，便于 RAG 测试。

## 建议测试问题

| 问题 | 期望答案要点 |
|------|----------------|
| 企业版多少钱？ | {FACTS['enterprise_price']} 元/年 |
| 免费版能建几个智能体？ | {FACTS['max_agents_free']} 个 |
| 专业版知识库多大？ | {FACTS['kb_storage_gb_pro']} GB |
| 退款几天内？ | {FACTS['refund_days']} 天 |
| 客服邮箱？ | {FACTS['support_email']} |
| SLA 是多少？ | {FACTS['sla_uptime']} |

## 页面上传说明

当前知识库页「上传文档」支持：**TXT、MD、JSON、CSV、LOG、PDF、DOCX**。

Excel（`.xlsx`）与 PPT（`.pptx`）已生成，后端 Tika 可解析；若页面上传被限制，可用 API 或后续扩展 accept 类型。

## 重新生成

```bash
python test-fixtures/knowledge-base/generate_samples.py
```
""",
        encoding="utf-8",
    )


def main() -> None:
    write_text_files()
    write_docx()
    write_xlsx()
    write_pptx()
    write_pdf()
    write_readme()
    print(f"Generated samples in {OUT}")


if __name__ == "__main__":
    main()
